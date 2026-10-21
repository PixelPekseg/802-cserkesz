#!/usr/bin/env python3
"""A Pénzügyek oldal (penzugyek.html) Tagdíj / Nagytábordíj / Kcs tábordíj füleit frissíti
fülenként egy-egy Google Táblázatból.

Minden fülnek külön táblázata van, ugyanazokkal a lapokkal:
  Alapadatok   – két oszlop (Mező · Érték): Megjelenés kezdete, Megjelenés vége, Cím,
                 Bevezető szöveg, Kiemelt mondat, Alapüzenet, ECSET kódok linkje,
                 Támogatás űrlap linkje, Költségek címe / bevezetője / megjegyzése
  Díjak        – Kategória · Összeg · Részletezés · Magyarázat címe · Magyarázat · Megjelenjen?
                 (a díjkártyák és alattuk a magyarázó bekezdések; egy sorban lehet díj és
                 magyarázat is, egymástól függetlenül. A magyarázat címe félkövéren a bekezdés elejére kerül.)
  Részletek    – Cím · Befizetési határidő · Összeg · Megjegyzés · Megjelenjen?
  Költségek    – Tétel · Összeg (Ft) · Megjelenjen?   (üresen hagyva a rész nem jelenik meg)
A táblázatot egészben, Microsoft Excel (.xlsx) formátumban kell közzétenni a weben, így
fülenként egyetlen cím kell (PENZUGY_TAGDIJ_URL, PENZUGY_NAGYTABOR_URL, PENZUGY_KCSTABOR_URL).

A fül tartalma a Megjelenés kezdete és vége között látszik (a vég napja még beleszámít; üres
vég = nincs vége); azon kívül, vagy ha nincs kezdet, csak az Alapüzenet. Az időszakot a
script.js (updateFees) minden megnyitáskor újra ellenőrzi, így nem kell hozzá újabb futás.
A banki adatok, a közlemény-figyelmeztetés és az ECSET doboz szövege minden fülön azonos, a
kódban rögzített (a számlaszámot szándékosan nem lehet a táblázatból átírni).

Szövegcellákban: új sor = új bekezdés, **szöveg** = félkövér.

Használat:
    python tools/sync_penzugyek.py                       # a három cím környezeti változókból
    python tools/sync_penzugyek.py --tagdij-file t.xlsx  # helyi fájlból (ugyanígy --nagytabor-file, --kcstabor-file)

A penzugyek.html <!-- PENZUGY_<FÜL>:START/END --> közötti részét írja újra. Ha egy fülnek nincs
címe, azt kihagyja (marad a régi). Hibás táblázatnál az adott fül változatlan marad, a többi
frissül, a futás pedig hibával zárul.
"""
import argparse
import datetime
import html
import io
import json
import os
import re
import sys
import urllib.request

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import sync_common as common

PAGE = os.path.join(common.ROOT, 'penzugyek.html')
TABS = [
    # (kulcs, környezeti változó, jelölő)
    ('tagdij', 'PENZUGY_TAGDIJ_URL', 'PENZUGY_TAGDIJ'),
    ('nagytabor', 'PENZUGY_NAGYTABOR_URL', 'PENZUGY_NAGYTABOR'),
    ('kcstabor', 'PENZUGY_KCSTABOR_URL', 'PENZUGY_KCSTABOR'),
]
# a fül neve a címsávban (a script.js a data-label-t írja a h1-be)
LABELS = {'tagdij': 'Tagdíj', 'nagytabor': 'Nagytábordíj', 'kcstabor': 'Kcs tábordíj'}
INDENT = ' ' * 6

# Az Alapadatok lap mezői (a Mező oszlop ELEJE alapján, kisbetűsen)
FIELDS = {
    'from': 'megjelenés kezdete', 'until': 'megjelenés vége', 'title': 'cím',
    'intro': 'bevezető', 'key': 'kiemelt', 'empty': 'alapüzenet',
    'ecset': 'ecset', 'form': 'támogatás', 'costs_title': 'költségek címe',
    'costs_intro': 'költségek bevezető', 'costs_note': 'költségek megjegyzés',
}
SHEETS = {
    'prices': ('Díjak', {'name': 'kategória', 'amount': 'összeg', 'split': 'részletezés',
                         'rule_title': 'magyarázat címe', 'rule_text': 'magyarázat', 'visible': 'megjelenjen'}),
    'items': ('Részletek', {'title': 'cím', 'deadline': 'befizetési', 'amount': 'összeg', 'note': 'megjegyzés',
                            'visible': 'megjelenjen'}),
    'costs': ('Költségek', {'name': 'tétel', 'amount': 'összeg', 'visible': 'megjelenjen'}),
}
DEFAULT_EMPTY = {'tagdij': 'Jelenleg nincs tagdíj befizetési időszak!'}
MONTHS = ['január', 'február', 'március', 'április', 'május', 'június', 'július',
          'augusztus', 'szeptember', 'október', 'november', 'december']

COPY_ICON = ('<svg class="icon-copy" viewBox="0 0 24 24" aria-hidden="true"><rect x="9" y="9" width="11" height="11" rx="2"/>'
             '<path d="M5 15V6a2 2 0 0 1 2-2h9"/></svg><svg class="icon-check" viewBox="0 0 24 24" aria-hidden="true">'
             '<path d="M5 13l4 4L19 7"/></svg>')


class SheetError(Exception):
    pass


# ---------- beolvasás ----------

def cell_text(value):
    """Cellaérték szövegként (a dátumokat ÉÉÉÉ.HH.NN. alakra, az egész számokat tizedes nélkül)."""
    if value is None:
        return ''
    if isinstance(value, datetime.datetime):
        value = value.date()
    if isinstance(value, datetime.date):
        return value.strftime('%Y.%m.%d.')
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value).replace('\r\n', '\n').replace('\r', '\n').strip()
    return '' if common.is_sheet_error(text) else text


def load_workbook(data):
    try:
        import openpyxl
    except ImportError:
        common.fail('Hiányzik az openpyxl csomag (pip install openpyxl).')
    try:
        return openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    except Exception as error:
        raise SheetError(f'a letöltött fájl nem olvasható Excel-táblázatként ({error}). '
                         'Biztosan a TELJES dokumentum van közzétéve, Microsoft Excel (.xlsx) formátumban?')


def find_sheet(book, name):
    for sheet in book.worksheets:
        if sheet.title.strip().lower() == name.lower():
            return sheet
    return None


def read_fields(book):
    sheet = find_sheet(book, 'Alapadatok')
    if sheet is None:
        raise SheetError('nincs "Alapadatok" lap.')
    fields = {key: '' for key in FIELDS}
    for row in sheet.iter_rows(values_only=True):
        if not row or row[0] is None:
            continue
        label = cell_text(row[0]).lower()
        value = cell_text(row[1]) if len(row) > 1 else ''
        # a hosszabb előtag nyer (pl. "költségek címe" ne a "cím"-hez kerüljön)
        matches = [key for key, prefix in FIELDS.items() if label.startswith(prefix)]
        if matches:
            key = max(matches, key=lambda k: len(FIELDS[k]))
            fields[key] = value
    return fields


def read_rows(book, sheet_name, columns):
    """Egy lap sorai szótárként; a "Megjelenjen?" = nem sorokat kihagyja. Hiányzó lap = üres lista."""
    sheet = find_sheet(book, sheet_name)
    if sheet is None:
        common.warn(f'Nincs "{sheet_name}" lap, ezt a részt kihagyom.')
        return []
    rows = [[cell_text(v) for v in row] for row in sheet.iter_rows(values_only=True)]
    if not rows:
        return []
    header = [c.lower() for c in rows[0]]
    index = {}
    # a hosszabb előtag előbb választ (pl. "magyarázat címe" ne a "magyarázat"-hoz kerüljön)
    for key, prefix in sorted(columns.items(), key=lambda item: -len(item[1])):
        for i, cell in enumerate(header):
            if cell.startswith(prefix) and i not in index.values():
                index[key] = i
                break
    missing = [columns[k] for k in columns if k not in index and k != 'visible']
    if missing:
        raise SheetError(f'a "{sheet_name}" lap fejlécéből hiányzik: {", ".join(missing)}.')
    result = []
    for row in rows[1:]:
        record = {key: (row[i] if i < len(row) else '') for key, i in index.items()}
        if not any(v for k, v in record.items() if k != 'visible'):
            continue
        if record.get('visible', '').lower() == 'nem':
            continue
        result.append(record)
    return result


def parse_date(text):
    match = re.search(r'(\d{4})\D+(\d{1,2})\D+(\d{1,2})', text or '')
    if not match:
        return None
    try:
        return datetime.date(*map(int, match.groups()))
    except ValueError:
        return None


def nice_date(text):
    """2026.11.30. -> 2026. november 30. (ha nem dátum, változatlanul hagyja)."""
    day = parse_date(text)
    return f'{day.year}. {MONTHS[day.month - 1]} {day.day}.' if day else text


def money(text):
    """Puszta szám -> "15 000 Ft"; minden más beírás marad."""
    digits = text.replace(' ', '').replace(' ', '')
    if re.fullmatch(r'\d+', digits):
        return f'{int(digits):,}'.replace(',', ' ') + ' Ft'
    return text


def check_link(text, what):
    if text and not re.match(r'https://', text):
        common.warn(f'A(z) {what} nem https:// kezdetű, kihagyom: "{text}"')
        return ''
    return text


# ---------- HTML ----------

def esc(text):
    return html.escape(text, quote=True)


def rich(text):
    """Escape-elt szöveg, **félkövér** jelöléssel."""
    return re.sub(r'\*\*(.+?)\*\*', r'<strong>\1</strong>', esc(text))


def paragraphs(text):
    return [rich(p.strip()) for p in text.split('\n') if p.strip()]


def copy_value(label, value, copy, mono=False, aria=None):
    aria = aria or label
    cls = 'bank-value bank-number' if mono else 'bank-value'
    return (f'              <p class="bank-label">{label}</p>\n'
            f'              <div class="bank-copy">\n'
            f'                <p class="{cls}">{esc(value)}</p>\n'
            f'                <button class="copy-btn" type="button" data-copy="{esc(copy)}" data-copy-label="{esc(aria)}" '
            f'aria-label="{esc(aria)} másolása" title="Másolás">{COPY_ICON}</button>\n'
            f'              </div>')


def render_pay(ecset_link):
    """Banki adatok + közlemény-figyelmeztetés + ECSET doboz: minden fülön azonos."""
    ecset = ''
    if ecset_link:
        ecset = f'''

            <div class="fee-box fee-ecset">
              <div>
                <h3>ECSET kód</h3>
                <p>A gyermek ECSET kódját a szülőkkel megosztott névsorban találja.</p>
                <p class="fee-box-note">
                  A fájlt csak akkor tudja megnyitni, ha azzal az e-mail-címmel van bejelentkezve
                  a Google-fiókjába, amelyik a csapat levelezőlistáján is szerepel.
                </p>
              </div>
              <a class="btn btn-primary" href="{esc(ecset_link)}" target="_blank" rel="noopener noreferrer">ECSET kódok megnyitása</a>
            </div>'''
    return f'''          <div class="fee-pay">
            <div class="bank-card">
              <h3>Banki átutalás az egyesület számlájára</h3>
{copy_value('Kedvezményezett neve', 'Gödöllői Cserkészetért Egyesület', 'Gödöllői Cserkészetért Egyesület')}
{copy_value('Számlaszám', '1620 0223 1007 7922 0000 0000', '162002231007792200000000', mono=True)}
{copy_value('vagy másodlagos banki azonosító', '+36204924970', '+36204924970', mono=True, aria='Másodlagos banki azonosító')}
              <p class="bank-label">Közlemény</p>
              <p class="bank-value">„Adomány” + ECSET kód</p>
            </div>

            <div class="fee-box fee-warning">
              <h3>Fontos: csak ez a két adat legyen a közleményben!</h3>
              <p>
                A közleménybe az „Adomány” szót és a gyermek ECSET kódját írja, ezeken kívül
                semmilyen más adatot ne adjon meg.
              </p>
              <p class="fee-example">Példa a közleményre: <strong>Adomány AAAAAA-1</strong></p>
              <ul>
                <li>Ha hiányzik az „Adomány” szó, a pénzt visszautaljuk, mert csak adományt fogadhatunk.</li>
                <li>Ha az ECSET kód hiányzik vagy nem a megfelelő formában szerepel, a rendszer nem ismeri fel a befizetést, és a pénzt szintén visszautaljuk.</li>
                <li>A visszautalás miatt mindkét esetben 1 000 Ft plusz költséget számolunk fel.</li>
              </ul>
            </div>{ecset}
          </div>
'''


def render_form(form_link):
    """Támogatás igénylése: beágyazott űrlap; a src-t a script.js csak akkor tölti be, ha a fül látszik."""
    base = re.sub(r'\?.*$', '', form_link)
    embed = base + '?embedded=true' if base.endswith('/viewform') else form_link
    return f'''          <div class="fee-form">
            <h3>Támogatás igénylése</h3>
            <p>
              A csapat minden évben támogatást oszt a tábordíjhoz és a tagdíjhoz az arra jelentkező
              családoknak. Amennyiben a befizetés megterhelő Önöknek, kérjük, nyugodtan jelezzék
              az alábbi űrlap kitöltésével.
            </p>
            <div class="fee-form-frame">
              <iframe data-src="{esc(embed)}" title="Támogatás igénylése (Google Űrlap)" loading="lazy">Betöltés…</iframe>
            </div>
            <p class="fee-form-link">Nem jelenik meg az űrlap? <a class="inline-link" href="{esc(form_link)}" target="_blank" rel="noopener noreferrer">Megnyitás új lapon</a></p>
          </div>
'''


def render_tab(key, data, today):
    f = data['fields']
    empty = f['empty'] or DEFAULT_EMPTY.get(key, 'Jelenleg nincs tábordíj befizetési időszak!')
    start, end = parse_date(f['from']), parse_date(f['until'])
    if f['from'] and not start:
        common.warn(f'{key}: a Megjelenés kezdete nem dátum ("{f["from"]}"), csak az alapüzenet jelenik meg.')
    if f['until'] and not end:
        common.warn(f'{key}: a Megjelenés vége nem dátum ("{f["until"]}"), nincs vége az időszaknak.')
    has_content = bool(start and f['title'])
    window = ''
    if has_content:
        window = f' data-from="{start.isoformat()}"' + (f' data-until="{end.isoformat()}"' if end else '')
    active = has_content and start <= today and (not end or today <= end)

    out = [f'<div class="fee-panel" id="panel-{key}" data-tab="{key}" data-label="{LABELS[key]}"{window}>']
    if not has_content:
        out.append('        <div class="fee-content" hidden></div>')
    else:
        out.append('        <div class="fee-content"' + ('' if active else ' hidden') + '>')
        # 1. cím, bevezető, kiemelt mondat
        out.append('          <div class="fee-intro">')
        out.append(f'            <h2>{esc(f["title"])}</h2>')
        out += [f'            <p>{p}</p>' for p in paragraphs(f['intro'])]
        out.append('          </div>\n')
        if f['key']:
            out.append(f'          <p class="fee-key">{rich(f["key"])}</p>\n')
        # 2. díjkártyák
        prices = [row for row in data['prices'] if row.get('name')]
        rules = [row for row in data['prices'] if row.get('rule_text')]
        if prices:
            out.append('          <ul class="price-grid">')
            for row in prices:
                split = f'<span class="price-split">{esc(money(row.get("split", "")))}</span>' if row.get('split') else ''
                out.append(f'            <li><span class="price-name">{esc(row["name"])}</span>'
                           f'<span class="price-amount">{esc(money(row["amount"]))}</span>{split}</li>')
            out.append('          </ul>\n')
        # 3. díjkártya-magyarázatok
        if rules:
            out.append('          <div class="fee-rules">')
            for row in rules:
                for i, p in enumerate(paragraphs(row['rule_text'])):
                    title = f'<strong>{esc(row["rule_title"].rstrip(":"))}:</strong> ' if row.get('rule_title') and i == 0 else ''
                    out.append(f'            <p>{title}{p}</p>')
            out.append('          </div>\n')
        # 4. részletek
        if data['items']:
            out.append('          <div class="fee-items">')
            for row in data['items']:
                note = f'\n              <p class="fee-item-note">{rich(row["note"])}</p>' if row.get('note') else ''
                out.append(f'''            <div class="fee-item">
              <h3>{esc(row["title"])}</h3>
              <dl class="fee-facts">
                <div><dt>Befizetési határidő</dt><dd>{esc(nice_date(row.get("deadline", "")))}</dd></div>
                <div><dt>Összeg</dt><dd>{esc(money(row.get("amount", "")))}</dd></div>
              </dl>{note}
            </div>''')
            out.append('          </div>\n')
        # 5. fix kártyák
        out.append(render_pay(f['ecset']))
        # 6. költségek
        costs = []
        for row in data['costs']:
            digits = row['amount'].replace(' ', '').replace(' ', '').removesuffix('Ft')
            if not digits.isdigit():
                common.warn(f'{key}: a költség összege nem szám, kihagyom: "{row["name"]}" – "{row["amount"]}"')
                continue
            costs.append((row['name'], int(digits)))
        total = sum(v for _, v in costs)
        if costs and total:
            out.append('          <div class="fee-costs">')
            out.append(f'            <h3>{esc(f["costs_title"] or "Mire fordítjuk a befizetéseket?")}</h3>')
            out += [f'            <p>{p}</p>' for p in paragraphs(f['costs_intro'])]
            out.append('            <ul class="cost-list">')
            for name, value in costs:
                out.append(f'''              <li style="--share:{round(value / total * 100)}%">
                <span class="cost-name">{esc(name)}</span>
                <span class="cost-amount">{money(str(value))}</span>
                <span class="cost-bar" aria-hidden="true"></span>
              </li>''')
            out.append('            </ul>')
            out.append(f'            <p class="cost-total">Összesen évente: <strong>{money(str(total))}</strong></p>')
            out += [f'            <p class="cost-note">{p}</p>' for p in paragraphs(f['costs_note'])]
            out.append('          </div>\n')
        # 7. támogatás igénylése
        if f['form']:
            out.append(render_form(f['form']))
        out[-1] = out[-1].rstrip('\n')
        out.append('        </div>')
    out.append(f'        <p class="fee-empty"' + (' hidden' if active else '') + f'>{esc(empty)}</p>')
    out.append('      </div>')
    return INDENT + '\n'.join(out)


# ---------- futtatás ----------

def fetch_bytes(url):
    try:
        with urllib.request.urlopen(url, timeout=30) as response:
            return response.read()
    except Exception as error:
        raise SheetError(f'a letöltés nem sikerült: {error}')


def read_tab(data):
    book = load_workbook(data)
    fields = read_fields(book)
    fields['ecset'] = check_link(fields['ecset'], 'ECSET kódok linkje')
    fields['form'] = check_link(fields['form'], 'támogatás űrlap linkje')
    result = {'fields': fields}
    for key, (sheet, columns) in SHEETS.items():
        result[key] = read_rows(book, sheet, columns)
    return result


def main():
    parser = argparse.ArgumentParser()
    for key, _, _ in TABS:
        parser.add_argument(f'--{key}-file')
    parser.add_argument('--today', help='teszteléshez: ÉÉÉÉ-HH-NN')
    args = parser.parse_args()
    today = datetime.date.fromisoformat(args.today) if args.today else datetime.date.today()

    updated, errors = [], []
    for key, env_name, marker in TABS:
        path = getattr(args, f'{key}_file')
        url = os.environ.get(env_name)
        if not path and not url:
            common.warn(f'Nincs megadva a(z) {key} táblázat címe ({env_name}), ezt a fület kihagyom.')
            continue
        try:
            raw = open(path, 'rb').read() if path else fetch_bytes(url)
            data = read_tab(raw)
            body = render_tab(key, data, today)
        except SheetError as error:
            errors.append(f'{key}: {error}')
            continue
        if common.replace_block(PAGE, marker, body, INDENT):
            updated.append(key)
        common.save_snapshot(os.path.join(common.ROOT, 'data', f'penzugy_{key}.json'),
                             json.dumps(data, ensure_ascii=False, indent=2) + '\n')

    print(f'Frissített fülek: {", ".join(updated) or "nincs"}.')
    if errors:
        common.fail('Hibás táblázat, ez a fül változatlan maradt – ' + ' | '.join(errors))


if __name__ == '__main__':
    main()
